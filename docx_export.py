
# docx_export.py
# Adds a brand-new /export/word route producing an executive Word report.
# This does not touch any existing route in app.py — it's a new endpoint,
# registered the same way app.py registers its own routes, just from this
# file instead. app.py only needs one more import line to load it.
#
# Reuses the same thread-safe, N+1-free chart/data approach as pdf_export.py:
# matplotlib configured once at import time (with the same font-cache fix),
# and all agent data pulled from the two bulk get_all_ot()/get_all_perf()
# calls, grouped locally — never per-agent Firebase round-trips.

import os
import tempfile

_mpl_cache_dir = os.path.join(tempfile.gettempdir(), "mpl_cache_team_shine")
os.makedirs(_mpl_cache_dir, exist_ok=True)
os.environ.setdefault("MPLCONFIGDIR", _mpl_cache_dir)

import matplotlib
matplotlib.use("Agg")
matplotlib.rcParams["font.family"] = "DejaVu Sans"
from matplotlib.figure import Figure
from matplotlib.backends.backend_agg import FigureCanvasAgg

from io import BytesIO
from datetime import datetime
from collections import defaultdict
import calendar
import traceback

from flask import Response

from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

from app import (
    app, login_required, page, PH_TZ,
    get_all, get_all_ot, get_all_perf, calc, calc_perf, parse_date,
    get_all_announcements, get_announcement_reads,
)

try:
    import coaching
    HAS_COACHING = True
except Exception:
    HAS_COACHING = False

NAVY = RGBColor(0x0B, 0x11, 0x20)
GOLD = RGBColor(0xFB, 0xBF, 0x24)
GREEN = RGBColor(0x16, 0xA3, 0x4A)
RED = RGBColor(0xDC, 0x26, 0x26)
AMBER = RGBColor(0xD9, 0x77, 0x06)
GREY = RGBColor(0x64, 0x74, 0x8B)


# ---------- Chart helper (thread-safe: Figure + FigureCanvasAgg, no pyplot) ----------

def _fig_to_png_bytes(fig):
    fig.set_dpi(150)
    canvas_ = FigureCanvasAgg(fig)
    buf = BytesIO()
    canvas_.print_png(buf)
    buf.seek(0)
    return buf


def make_ot_trend_chart(months, normal_m, restday_m, total_m, loss_m):
    fig = Figure(figsize=(7.2, 3.0))
    ax = fig.add_subplot(111)
    ax.plot(months, normal_m, color="#22c55e", marker="o", label="Normal OT", linewidth=1.8)
    ax.plot(months, restday_m, color="#3b82f6", marker="o", label="Restday OT", linewidth=1.8)
    ax.plot(months, total_m, color="#fbbf24", marker="o", label="Total OT", linewidth=2.2)
    ax.plot(months, loss_m, color="#ef4444", marker="o", label="Loss", linewidth=1.8)
    ax.set_title("Monthly OT & Loss Trend", fontsize=11, fontweight="bold")
    ax.legend(fontsize=7, ncol=4, loc="upper center", bbox_to_anchor=(0.5, -0.15))
    ax.grid(alpha=0.25)
    ax.tick_params(labelsize=8)
    fig.tight_layout()
    return _fig_to_png_bytes(fig)


def make_quality_trend_chart(months, qa_m, aht_m):
    fig = Figure(figsize=(7.2, 3.0))
    ax = fig.add_subplot(111)
    ax.plot(months, qa_m, color="#8b5cf6", marker="o", label="QA %", linewidth=1.8)
    ax2 = ax.twinx()
    ax2.plot(months, aht_m, color="#f97316", marker="s", linestyle="--", label="AHT (min)", linewidth=1.5)
    ax.set_title("Monthly Quality Trend (QA vs AHT)", fontsize=11, fontweight="bold")
    lines1, labels1 = ax.get_legend_handles_labels()
    lines2, labels2 = ax2.get_legend_handles_labels()
    ax.legend(lines1 + lines2, labels1 + labels2, fontsize=7, ncol=2, loc="upper center", bbox_to_anchor=(0.5, -0.15))
    ax.grid(alpha=0.25)
    ax.tick_params(labelsize=8)
    ax2.tick_params(labelsize=8)
    fig.tight_layout()
    return _fig_to_png_bytes(fig)


# ---------- Word styling helpers ----------

def _shade_cell(cell, hex_color):
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_color)
    cell._tc.get_or_add_tcPr().append(shd)


def _set_cell_text(cell, text, bold=False, color=None, size=9, align=None):
    cell.text = ""
    p = cell.paragraphs[0]
    if align:
        p.alignment = align
    run = p.add_run(str(text))
    run.bold = bold
    run.font.size = Pt(size)
    if color:
        run.font.color.rgb = color


def add_heading_band(doc, text, subtitle=None):
    table = doc.add_table(rows=2 if subtitle else 1, cols=1)
    table.autofit = True
    cell = table.rows[0].cells[0]
    _shade_cell(cell, "0B1120")
    _set_cell_text(cell, text, bold=True, color=RGBColor(0xFF, 0xFF, 0xFF), size=18)
    if subtitle:
        cell2 = table.rows[1].cells[0]
        _shade_cell(cell2, "0B1120")
        _set_cell_text(cell2, subtitle, bold=False, color=RGBColor(0xCB, 0xD5, 0xE1), size=9)


def add_section_heading(doc, text):
    h = doc.add_heading(text, level=2)
    for run in h.runs:
        run.font.color.rgb = NAVY


def build_kpi_table(doc, kpi_rows):
    table = doc.add_table(rows=1, cols=3)
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    hdr = table.rows[0].cells
    for i, h in enumerate(["Metric", "Value", "Status"]):
        _shade_cell(hdr[i], "0B1120")
        _set_cell_text(hdr[i], h, bold=True, color=GOLD, size=10)
    for i, (metric, value, status) in enumerate(kpi_rows):
        row = table.add_row().cells
        if i % 2 == 1:
            for c in row:
                _shade_cell(c, "F1F5F9")
        _set_cell_text(row[0], metric, size=9)
        _set_cell_text(row[1], value, size=9)
        sc = GREEN if status == "Good" else RED if status == "Attention" else AMBER if status == "Watch" else GREY
        _set_cell_text(row[2], status, bold=True, color=sc, size=9)
    return table


def build_agent_table(doc, agent_stats):
    headers = ["Name", "Tencent ID", "OT (h)", "Loss (h)", "QA %", "AHT (m)", "WH Comp."]
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    hdr = table.rows[0].cells
    for i, h in enumerate(headers):
        _shade_cell(hdr[i], "0B1120")
        _set_cell_text(hdr[i], h, bold=True, color=GOLD, size=8.5)
    ranked = sorted(agent_stats, key=lambda x: -x["wh_comp"])
    for i, a in enumerate(ranked):
        row = table.add_row().cells
        if i % 2 == 1:
            for c in row:
                _shade_cell(c, "F1F5F9")
        _set_cell_text(row[0], a["name"], size=8.5)
        _set_cell_text(row[1], a["tid"], size=8.5)
        _set_cell_text(row[2], round(a["tot"], 1), size=8.5)
        loss_color = RED if a["loss"] >= 4 else None
        _set_cell_text(row[3], round(a["loss"], 1), bold=bool(loss_color), color=loss_color, size=8.5)
        _set_cell_text(row[4], round(a["qa"], 1) if a["qa"] else "-", size=8.5)
        _set_cell_text(row[5], round(a["aht"], 1) if a["aht"] else "-", size=8.5)
        _set_cell_text(row[6], f"{round(a['wh_comp'],1)}%", size=8.5)
    return table


def build_risk_table(doc, at_risk):
    headers = ["Name", "Tencent ID", "Loss (h)", "Net OT (h)", "QA %"]
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    hdr = table.rows[0].cells
    for i, h in enumerate(headers):
        _shade_cell(hdr[i], "DC2626")
        _set_cell_text(hdr[i], h, bold=True, color=RGBColor(0xFF, 0xFF, 0xFF), size=8.5)
    for i, a in enumerate(at_risk[:15]):
        row = table.add_row().cells
        if i % 2 == 1:
            for c in row:
                _shade_cell(c, "FEF2F2")
        _set_cell_text(row[0], a["name"], size=8.5)
        _set_cell_text(row[1], a["tid"], size=8.5)
        _set_cell_text(row[2], round(a["loss"], 1), size=8.5)
        _set_cell_text(row[3], round(a["net"], 1), size=8.5)
        _set_cell_text(row[4], round(a["qa"], 1) if a["qa"] else "-", size=8.5)
    return table


# ---------- Report builder ----------

def build_docx_report():
    agents = get_all()
    ot_logs = get_all_ot()
    perf_logs = get_all_perf()
    anns = get_all_announcements()
    reads = get_announcement_reads()

    ot_by_agent = defaultdict(list)
    for l in ot_logs:
        ot_by_agent[str(l.get("agent_id"))].append(l)
    perf_by_agent = defaultdict(list)
    for l in perf_logs:
        perf_by_agent[str(l.get("agent_id"))].append(l)

    agent_stats = []
    team_n = team_r = team_loss = team_tot = 0.0
    for a in agents:
        logs = ot_by_agent.get(str(a.get("id")), [])
        n, r, tot, loss, net = calc(logs)
        team_n += n; team_r += r; team_loss += loss; team_tot += tot
        plogs = perf_by_agent.get(str(a.get("id")), [])
        la, lq, lc, lf, aa, qa, ac, af = calc_perf(plogs)
        wh_target = float(a.get("WORKING_HOURS_TARGET", 220))
        wh_actual = wh_target - loss
        wh_comp = (wh_actual / wh_target * 100) if wh_target > 0 else 0
        agent_stats.append({
            "name": a.get("NAME", ""), "tid": a.get("TENCENT_ID", ""),
            "tot": tot, "loss": loss, "net": net,
            "aht": la, "qa": lq,
            "wh_target": wh_target, "wh_actual": wh_actual, "wh_comp": wh_comp,
        })

    team_qa = [a["qa"] for a in agent_stats if a["qa"] > 0]
    team_aht = [a["aht"] for a in agent_stats if a["aht"] > 0]
    total_wh = sum(a["wh_target"] for a in agent_stats)
    total_actual = total_wh - team_loss
    overall_att = (total_actual / total_wh * 100) if total_wh > 0 else 0
    avg_qa = round(sum(team_qa) / len(team_qa), 1) if team_qa else 0
    avg_aht = round(sum(team_aht) / len(team_aht), 1) if team_aht else 0
    at_risk = sorted([a for a in agent_stats if a["loss"] >= 4], key=lambda x: -x["loss"])

    coaching_line = ""
    if HAS_COACHING:
        try:
            all_coaching = coaching.get_all_coaching()
            c_reads = coaching.get_coaching_reads()
            confirmed_ids = {r.get("coaching_id") for r in c_reads}
            pending = [c for c in all_coaching if c.get("id") not in confirmed_ids]
            coaching_line = f" {len(all_coaching)} coaching sessions logged this period, {len(pending)} pending agent confirmation."
        except Exception:
            coaching_line = ""

    m_ot = defaultdict(list)
    for l in ot_logs:
        dt = parse_date(l.get("date", ""))
        if dt:
            m_ot[dt.month].append(l)
    m_perf = defaultdict(list)
    for l in perf_logs:
        dt = parse_date(l.get("date", ""))
        if dt:
            m_perf[dt.month].append(l)

    months_lbl = [calendar.month_abbr[m] for m in range(1, 13)]
    normal_m, restday_m, total_m, loss_m = [], [], [], []
    qa_m, aht_m = [], []
    for m in range(1, 13):
        n, r, tot, loss, net = calc(m_ot.get(m, []))
        normal_m.append(round(n, 1)); restday_m.append(round(r, 1))
        total_m.append(round(tot, 1)); loss_m.append(round(loss, 1))
        la, lq, lc, lf, aa, qav, acv, afv = calc_perf(m_perf.get(m, []))
        qa_m.append(round(qav, 1)); aht_m.append(round(aa, 1))

    # ---- Build the document ----
    doc = Document()
    section = doc.sections[0]
    section.left_margin = Inches(0.6)
    section.right_margin = Inches(0.6)

    add_heading_band(
        doc, "TEAM SHINE M9",
        f"Executive Performance Report  •  Generated {datetime.now(PH_TZ).strftime('%B %d, %Y — %I:%M %p')}  •  Team size: {len(agents)} agents",
    )
    doc.add_paragraph()

    add_section_heading(doc, "Executive Summary")
    att_word = "strong" if overall_att >= 95 else "adequate" if overall_att >= 90 else "a concern"
    qa_word = "on target" if avg_qa >= 80 or avg_qa == 0 else "below target"
    summary_text = (
        f"The team of {len(agents)} agents logged {round(team_tot,1)}h of overtime this period "
        f"against {round(team_loss,1)}h of loss, for a net contribution of {round(team_tot-team_loss,1)}h. "
        f"Overall attendance stands at {round(overall_att,1)}%, which is {att_word}. "
        f"Average QA is {avg_qa}% ({qa_word}) and average AHT is {avg_aht} min. "
        f"{len(at_risk)} agent(s) are flagged for loss hours of 4+ and may need follow-up.{coaching_line} "
        f"{len(anns)} announcement(s) were posted with {len(reads)} total confirmations."
    )
    p = doc.add_paragraph(summary_text)
    p.paragraph_format.space_after = Pt(10)

    add_section_heading(doc, "Team KPI Summary")
    kpi_rows = [
        ("Attendance", f"{round(overall_att,1)}%", "Good" if overall_att >= 95 else "Watch" if overall_att >= 90 else "Attention"),
        ("Team Avg QA", f"{avg_qa}%", "Good" if avg_qa >= 85 else "Watch" if avg_qa >= 75 else "Attention"),
        ("Team Avg AHT", f"{avg_aht} min", "—"),
        ("Total OT", f"{round(team_tot,1)}h", "—"),
        ("Total Loss", f"{round(team_loss,1)}h", "—"),
        ("Agents at Risk (Loss ≥4h)", str(len(at_risk)), "Attention" if at_risk else "Good"),
    ]
    build_kpi_table(doc, kpi_rows)
    doc.add_paragraph()

    try:
        add_section_heading(doc, "Trends")
        doc.add_picture(make_ot_trend_chart(months_lbl, normal_m, restday_m, total_m, loss_m), width=Inches(6.5))
        doc.add_picture(make_quality_trend_chart(months_lbl, qa_m, aht_m), width=Inches(6.5))
    except Exception:
        print("DOCX chart rendering failed:")
        traceback.print_exc()
        doc.add_paragraph("Charts could not be rendered for this report. The data tables below are unaffected.")

    if at_risk:
        add_section_heading(doc, "Agents Flagged for Follow-up (Loss ≥ 4h)")
        build_risk_table(doc, at_risk)
        doc.add_paragraph()

    doc.add_page_break()
    add_section_heading(doc, "Agent Detail")
    build_agent_table(doc, agent_stats)

    footer = doc.sections[0].footer
    fp = footer.paragraphs[0]
    fp.text = "TEAM SHINE M9 — Executive Performance Report  •  Confidential — Internal Use Only"
    fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for run in fp.runs:
        run.font.size = Pt(7)
        run.font.color.rgb = GREY

    buffer = BytesIO()
    doc.save(buffer)
    return buffer.getvalue()


# ---------- Flask view (new route — does not override anything existing) ----------

@app.route("/export/word")
@login_required
def export_word():
    try:
        docx_bytes = build_docx_report()
        return Response(
            docx_bytes,
            mimetype="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            headers={"Content-Disposition": "attachment;filename=Team_Shine_M9_Executive_Report.docx"},
        )
    except Exception:
        print("export_word error:")
        traceback.print_exc()
        return page(
            f"<div class='card-dark'><h6 style='color:#ef4444'>Word Export Error</h6>"
            f"<pre style='color:#fbbf24;font-size:10px'>{traceback.format_exc()}</pre>"
            f"<a href='/export' class='btn btn-sm btn-outline-light'>Back</a></div>"
        )
